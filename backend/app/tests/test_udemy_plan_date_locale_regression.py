from datetime import date
from io import BytesIO
import os

os.environ.setdefault('DATABASE_URL', 'sqlite+pysqlite:///:memory:')

from openpyxl import load_workbook

from app.services.academic.udemy_plan import UdemyPlanService


def test_generated_template_keeps_week_dates_as_ddmmyyyy_text():
    """Ambiguous dates must not be handed back to Excel as locale-sensitive date cells."""
    raw = UdemyPlanService.build_template()
    wb = load_workbook(BytesIO(raw), data_only=True)
    ws = wb['KeHoachUdemy']

    week_cells = [ws.cell(row=3, column=col) for col in range(7, 27, 2)]
    populated = [cell for cell in week_cells if cell.value not in (None, '')]
    assert populated
    assert all(isinstance(cell.value, str) for cell in populated)
    assert all(cell.number_format == '@' for cell in week_cells)
    assert ws['K3'].value == '03/10/2026'
    wb.close()


def test_strict_ddmmyyyy_parser_does_not_swap_day_and_month():
    class Cell:
        value = '03/10/2026'

    parsed = UdemyPlanService._parse_date(Cell())
    assert parsed == date(2026, 10, 3)
