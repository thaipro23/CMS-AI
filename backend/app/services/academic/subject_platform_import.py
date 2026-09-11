"""Validate complete Excel platform plans, then apply one term atomically."""
from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime, timedelta, timezone
from io import BytesIO
import hashlib
import re
import secrets
import zipfile
from typing import Any

from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.models.academic import AcademicSubject, AcademicSubjectDelivery, AcademicTerm
from app.services.academic.subject_delivery import AcademicSubjectDeliveryService
from app.services.object_storage import StorageError, get_object_storage


class SubjectPlatformImportService:
    MAX_UPLOAD_BYTES = 10 * 1024 * 1024
    MAX_ROWS = 10000
    PREVIEW_TTL = timedelta(hours=2)

    def __init__(self, db: Session):
        self.db = db

    def _scope(self, term_id: str, branch: str) -> str:
        branch = str(branch or '').strip().lower()
        term = self.db.get(AcademicTerm, term_id)
        if not term or not term.active:
            raise HTTPException(422, 'Học kỳ không tồn tại hoặc không còn hoạt động.')
        if branch not in {'poly', 'ptcd'} or str(term.branch or '').strip().lower() != branch:
            raise HTTPException(422, 'Học kỳ không thuộc hệ đã chọn.')
        return branch

    def _targets(self, term_id: str, branch: str, *, lock: bool = False):
        query = self.db.query(AcademicSubjectDelivery, AcademicSubject).join(
            AcademicSubject, AcademicSubject.id == AcademicSubjectDelivery.subject_id,
        ).filter(
            AcademicSubjectDelivery.term_id == term_id,
            func.lower(AcademicSubjectDelivery.branch) == branch,
            AcademicSubjectDelivery.active.is_(True),
            AcademicSubject.active.is_(True),
            func.lower(AcademicSubject.branch) == branch,
        ).order_by(AcademicSubjectDelivery.id)
        if lock:
            query = query.with_for_update(of=AcademicSubjectDelivery).populate_existing()
        by_code = defaultdict(list)
        for delivery, subject in query.all():
            by_code[str(subject.subject_code).strip().casefold()].append((delivery, subject))
        return by_code

    def preview(self, raw: bytes, *, term_id: str, branch: str, requested_by: str) -> dict[str, Any]:
        branch = self._scope(term_id, branch)
        if not raw or len(raw) > self.MAX_UPLOAD_BYTES:
            raise HTTPException(422, 'File Excel phải có dữ liệu và không vượt quá 10 MB.')
        book = None
        try:
            with zipfile.ZipFile(BytesIO(raw)) as archive:
                if len(archive.infolist()) > 5000 or sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                    raise HTTPException(422, 'File Excel vượt giới hạn dữ liệu giải nén.')
            book = load_workbook(BytesIO(raw), read_only=True, data_only=False)
            if len(book.worksheets) != 1:
                raise HTTPException(422, 'File kế hoạch phải có đúng một sheet dữ liệu.')
            sheet = book.worksheets[0]
            if sheet.max_column != 2 or sheet.max_row > self.MAX_ROWS + 1:
                raise HTTPException(422, 'File cần đúng hai cột Mã môn, Nền tảng và tối đa 10.000 dòng.')
            values = list(sheet.iter_rows(values_only=True))
            if not values or tuple(str(value or '').strip() for value in values[0]) != ('Mã môn', 'Nền tảng'):
                raise HTTPException(422, 'Hai cột phải theo đúng thứ tự: Mã môn, Nền tảng.')
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(422, 'Không đọc được file Excel. Hãy dùng file .xlsx hợp lệ.') from exc
        finally:
            if book:
                book.close()

        data = [(number, str(code or '').strip(), str(platform or '').strip())
                for number, (code, platform) in enumerate(values[1:], 2)
                if str(code or '').strip() or str(platform or '').strip()]
        if not data:
            raise HTTPException(422, 'File kế hoạch chưa có dòng dữ liệu.')
        counts = Counter(code.casefold() for _, code, _ in data if code)
        targets = self._targets(term_id, branch)
        rows = []
        for number, code, value in data:
            matches = targets.get(code.casefold(), [])
            platform = {'cms': 'cms', 'udemy': 'udemy', 'other': 'other', 'khác': 'other'}.get(value.casefold())
            status, message = 'matched', f'{len(matches)} Block'
            if code and counts[code.casefold()] > 1:
                status, message = 'duplicate', 'Mã môn bị trùng trong file (không phân biệt hoa/thường).'
            elif not code or platform is None:
                status, message = 'invalid', 'Cần Mã môn và nền tảng CMS, Udemy hoặc Khác (Other).'
            elif not matches:
                status, message = 'missing', 'Không tìm thấy mã môn trong học kỳ và hệ đã chọn.'
            elif len({subject.id for _, subject in matches}) != 1:
                status, message = 'duplicate', 'Mã môn khớp nhiều môn trong danh mục; cần sửa danh mục.'
            rows.append({
                'row_no': number, 'subject_code': code, 'subject_name': matches[0][1].subject_name if matches else '',
                'learning_platform': platform, 'status': status, 'message': message,
                'delivery_ids': sorted(delivery.id for delivery, _ in matches),
                'previous_platforms': {delivery.id: delivery.learning_platform for delivery, _ in matches},
            })
        summary = Counter(row['status'] for row in rows)
        return {
            'term_id': term_id, 'branch': branch, 'requested_by': requested_by,
            'created_at': datetime.now(timezone.utc).isoformat(), 'file_sha256': hashlib.sha256(raw).hexdigest(),
            'total_rows': len(rows), **{f'{status}_count': summary[status] for status in ['matched', 'missing', 'duplicate', 'invalid']},
            'can_apply': summary['matched'] == len(rows), 'rows': rows,
        }

    def persist_preview(self, preview: dict[str, Any]) -> str:
        token = secrets.token_hex(16)
        get_object_storage().put_json(f'subject-platform-previews/{token}.json', preview)
        return token

    def load_preview(self, token: str, *, requested_by: str) -> dict[str, Any]:
        if not re.fullmatch(r'[0-9a-f]{32}', token):
            raise HTTPException(422, 'Mã xem trước không hợp lệ.')
        try:
            preview = get_object_storage().read_json(f'subject-platform-previews/{token}.json')
        except StorageError as exc:
            raise HTTPException(404, 'Không tìm thấy bản xem trước. Hãy tải lại file.') from exc
        if preview.get('requested_by') != requested_by:
            raise HTTPException(403, 'Bản xem trước không thuộc người dùng hiện tại.')
        if datetime.now(timezone.utc) - datetime.fromisoformat(preview['created_at']) > self.PREVIEW_TTL:
            raise HTTPException(410, 'Bản xem trước hết hạn sau 2 giờ. Hãy tải lại file.')
        return preview

    def apply_preview(self, preview: dict[str, Any], *, actor: str) -> dict[str, Any]:
        if preview.get('requested_by') != actor:
            raise HTTPException(403, 'Bản xem trước không thuộc người dùng hiện tại.')
        if not preview.get('can_apply') or not preview.get('rows') or any(row['status'] != 'matched' for row in preview['rows']):
            raise HTTPException(422, 'Hãy sửa các dòng lỗi trước khi áp dụng kế hoạch.')
        branch = self._scope(preview['term_id'], preview['branch'])
        try:
            targets = self._targets(preview['term_id'], branch, lock=True)
            changes = []
            for row in preview['rows']:
                matches = targets.get(row['subject_code'].casefold(), [])
                if sorted(delivery.id for delivery, _ in matches) != row['delivery_ids']:
                    raise HTTPException(409, 'Danh mục môn đã thay đổi. Hãy xem trước file lại.')
                platform = AcademicSubjectDeliveryService.normalize_platform(row['learning_platform'])
                for delivery, _ in matches:
                    if delivery.learning_platform not in {row['previous_platforms'][delivery.id], platform}:
                        raise HTTPException(409, 'Nền tảng đã được cập nhật sau khi xem trước. Hãy xem trước file lại.')
                    if delivery.learning_platform != platform:
                        changes.append((delivery, platform, row['row_no']))
            now = datetime.utcnow()
            for delivery, platform, row_no in changes:
                metadata = dict(delivery.metadata_json or {})
                history = list(metadata.get('platform_history') or [])
                history.append({'from': delivery.learning_platform, 'to': platform, 'source': 'excel_plan_import',
                                'actor': actor, 'changed_at': now.isoformat(), 'source_row': row_no,
                                'file_sha256': preview['file_sha256']})
                metadata['platform_history'] = history[-100:]
                delivery.metadata_json = metadata
                delivery.learning_platform = platform
                delivery.configured_by = actor
                delivery.configured_at = now
                delivery.configuration_source = 'excel_plan_import'
            self.db.commit()
        except Exception:
            self.db.rollback()
            raise
        subjects = len(preview['rows'])
        return {'ok': True, 'subjects': subjects, 'updated': len(changes), 'message': f'Đã áp dụng kế hoạch cho {subjects} môn.'}
